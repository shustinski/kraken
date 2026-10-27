from datetime import UTC, datetime, timedelta
from uuid import uuid4
import csv

import pytest

from kraken_manager.application.work_plan_analytics import WorkPlanAnalytics
from kraken_manager.domain.common import DomainValidationError
from kraken_manager.domain.events import ActorSnapshot, EventEnvelope
from kraken_manager.domain.identity import Principal
from kraken_manager.domain.planning import WorkAssignment, WorkDependency, WorkEstimate, WorkItem, WorkPlan, WorkScope
from kraken_manager.infrastructure.reports import ReportService

BASE = datetime(2026, 10, 24, tzinfo=UTC)


def uid():
    return str(uuid4())


class History:
    def __init__(self):
        self.actor = Principal.local(subject="engineer", display_name="Engineer")
        self.plan = WorkPlan.create(uid(), "Restoration", at=BASE)
        self.events = []
        self.snapshot()

    def snapshot(self):
        self.events.append(
            EventEnvelope.create(
                stream_id=f"work-plan:{self.plan.id}",
                project_id=self.plan.project_id,
                revision=self.plan.revision + 1,
                event_type="work_plan.changed",
                payload=self.plan.to_dict(),
                actor=ActorSnapshot.from_principal(self.actor),
                recorded_at=self.plan.updated_at,
            )
        )

    def change(self, hour, method, *args, **kwargs):
        self.plan = getattr(self.plan, method)(
            *args, **kwargs, expected_revision=self.plan.revision, at=BASE + timedelta(hours=hour)
        )
        self.snapshot()
        return self.plan

    def report(self, hour=15, **kwargs):
        return WorkPlanAnalytics.report(
            self.events, project_id=self.plan.project_id, as_of=BASE + timedelta(hours=hour), start=BASE, **kwargs
        )


def populated():
    history = History()
    layer1, layer2, performer1, performer2 = uid(), uid(), uid(), uid()
    stage = WorkItem.create("Stage")
    first = WorkItem.create(
        "Trace",
        parent_id=stage.id,
        scope=WorkScope(layer_ids=(layer1,), frame_ids=(uid(),)),
        assignment=WorkAssignment(performer1),
        estimate=WorkEstimate(8, BASE, BASE + timedelta(hours=20)),
    )
    second = WorkItem.create(
        "Inspect",
        scope=WorkScope(layer_ids=(layer2,)),
        assignment=WorkAssignment(performer2),
        estimate=WorkEstimate(4, BASE, BASE + timedelta(hours=24)),
    )
    history.change(1, "add_item", stage)
    history.change(2, "add_item", first)
    history.change(3, "add_item", second)
    history.change(4, "capture_baseline")
    history.change(5, "revise_item", second.id, estimate=WorkEstimate(6, BASE, BASE + timedelta(hours=26)))
    history.change(6, "transition", "active")
    history.change(7, "move_item", first.id, "running")
    history.change(8, "record_progress", first.id, 1, 3)
    history.change(9, "move_item", first.id, "in_review")
    history.change(10, "move_item", first.id, "running")
    history.change(11, "record_progress", first.id, 1, 5)
    history.change(12, "move_item", first.id, "in_review")
    history.change(13, "move_item", first.id, "accepted")
    history.change(14, "move_item", second.id, "blocked", reason="Awaiting microscope")
    return history, first, second, performer1, performer2


def test_event_metrics_are_leaf_based_measured_and_do_not_duplicate_snapshot_acceptance():
    history, first, second, performer1, performer2 = populated()
    report = history.report()
    summary = report["summary"]
    assert summary["leaf_items"] == 2 and summary["accepted_items"] == 1
    assert summary["readiness_fraction"] == 0.5
    assert summary["reported_effort_hours"] == 5
    assert summary["mean_lead_hours"] == 11
    assert summary["mean_cycle_hours"] == 6
    assert summary["rework_fraction"] == 1 and summary["rework_attempts"] == 1
    assert summary["acceptance_events_in_period"] == 1
    assert summary["acceptances_per_calendar_day"] == pytest.approx(1.6)
    assert summary["accepted_frame_scope_count"] == 1
    assert summary["measured_compute_utilization"] is None
    assert report["blocking_reasons"] == {"Awaiting microscope": 1}
    assert report["forecast"]["execution_finish_at"] is None
    assert report["forecast"]["reason"] == "blocked_work"
    assert set(report["source_event_ids"]) == {e.event_id for e in history.events}
    row = next(r for r in report["items"] if r["item_id"] == first.id)
    assert row["hierarchy"] == ["Stage"]
    assert [datetime.fromisoformat(change["at"]).hour for change in row["changes"]] == [2, 7, 8, 9, 10, 11, 12, 13]
    assert row["event_ids"][-1] == history.events[-2].event_id


def test_historical_snapshot_uses_old_revision_and_never_leaks_future_acceptance():
    history, first, second, performer1, performer2 = populated()
    report = history.report(8)
    assert report["source_plan_revisions"] == {history.plan.id: 8}
    assert report["summary"]["accepted_items"] == 0
    assert report["summary"]["mean_lead_hours"] is None
    assert report["summary"]["reported_effort_hours"] == 3
    assert len(report["source_event_ids"]) == 9
    assert report["complete_event_history"][history.plan.id]
    assert report["baseline_variance"][0]["baseline_source_revision"] == 3


def test_baseline_variance_and_named_filter_metrics_have_defined_units():
    history, first, second, performer1, performer2 = populated()
    report = history.report(30)
    variances = {row["item_id"]: row for row in report["baseline_variance"]}
    assert variances[second.id]["planned_effort_delta_hours"] == 2
    assert variances[second.id]["planned_finish_delta_hours"] == 2
    assert variances[first.id]["acceptance_finish_variance_hours"] == -7
    assert variances[first.id]["baseline_plan_effort_hours"] == 12
    assert variances[first.id]["current_plan_effort_hours"] == 14
    assert report["summary"]["overdue_items"] == 1
    selected = history.report(performer_ids=[performer1])
    assert selected["summary"]["leaf_items"] == 1 and selected["summary"]["readiness_fraction"] == 1
    assert selected["items"][0]["item_id"] == first.id
    selected = history.report(states=["blocked"])
    assert selected["state_counts"] == {"blocked": 1}
    with pytest.raises(DomainValidationError, match="baseline"):
        history.report(baseline_ids={history.plan.id: uid()})


def test_forecast_uses_explicit_estimates_capacity_dependencies_and_calendar_units():
    history = History()
    performer1, performer2 = uid(), uid()
    first = WorkItem.create("First", assignment=WorkAssignment(performer1), estimate=WorkEstimate(8))
    second = WorkItem.create("Second", assignment=WorkAssignment(performer2), estimate=WorkEstimate(8))
    history.change(1, "add_item", first)
    history.change(2, "add_item", second)
    history.change(3, "add_dependency", WorkDependency(first.id, second.id))
    missing = history.report(4)
    assert missing["forecast"]["execution_finish_at"] is None
    assert missing["forecast"]["reason"] == "missing_capacity"
    report = history.report(4, capacity_by_performer={performer1: 8, performer2: 8})
    assert report["forecast"]["remaining_calendar_days"] == 2
    assert report["forecast"]["execution_finish_at"] == (BASE + timedelta(hours=4, days=2)).isoformat()
    assert report["forecast"]["acceptance_finish_at"] is None
    assert report["workload"][0]["reported_effort_hours"] == 0
    assert all(r["remaining_estimated_hours"] == 8 for r in report["workload"])
    overflow = history.report(4, capacity_by_performer={performer1: 1e-300, performer2: 1e-300})
    assert overflow["forecast"]["execution_finish_at"] is None


@pytest.mark.parametrize("capacity", [0, -1, True, float("inf"), float("nan")])
def test_invalid_capacity_is_rejected(capacity):
    history = History()
    with pytest.raises(DomainValidationError, match="capacity"):
        history.report(capacity_by_performer={uid(): capacity})


def test_missing_event_history_does_not_invent_lead_time_or_zero_estimates():
    history = History()
    item = WorkItem.create("Unestimated", assignment=WorkAssignment(uid()))
    history.change(1, "add_item", item)
    report = WorkPlanAnalytics.report(
        history.events[1:], project_id=history.plan.project_id, as_of=BASE + timedelta(hours=2)
    )
    assert not report["complete_event_history"][history.plan.id]
    assert report["items"][0]["first_observed_at"] is None
    assert report["items"][0]["lead_hours"] is None
    assert report["items"][0]["remaining_estimated_hours"] is None
    assert report["forecast"]["reason"] == "missing_estimate"


def test_report_exports_unknown_as_blank_and_defuses_formula_titles(tmp_path):
    from openpyxl import load_workbook

    history = History()
    history.change(1, "add_item", WorkItem.create('=HYPERLINK("unsafe")'))
    report = history.report(2)
    csv_path = ReportService().write_work_plan_csv(tmp_path / "plans.csv", report)
    with csv_path.open(encoding="utf-8", newline="") as stream:
        records = list(csv.DictReader(stream))
    assert any(row["title"].startswith("'=HYPERLINK") for row in records)
    unknown = next(row for row in records if row["section"] == "summary" and row["metric"] == "mean_lead_hours")
    assert unknown["value"] == ""
    assert {r["event_id"] for r in records if r["section"] == "provenance"} == set(report["source_event_ids"])
    path = ReportService().write_work_plan_xlsx(tmp_path / "plans.xlsx", report)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        assert {"Summary", "Tasks", "Forecast", "Provenance", "History"} <= set(workbook.sheetnames)
        headers = next(workbook["Tasks"].values)
        title_index = headers.index("title")
        assert next(iter(list(workbook["Tasks"].values)[1:]))[title_index].startswith("'=HYPERLINK")
        assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    finally:
        workbook.close()


def test_incomplete_history_does_not_fabricate_acceptance_event_time_or_throughput():
    history, first, second, performer1, performer2 = populated()
    # Remove the actual acceptance event; the later snapshot still carries the
    # accepted state but cannot prove when that transition occurred.
    source = [event for event in history.events if event.recorded_at != BASE + timedelta(hours=13)]
    report = WorkPlanAnalytics.report(source, project_id=history.plan.project_id, as_of=BASE + timedelta(hours=15))
    assert report["summary"]["accepted_items"] == 1
    assert report["summary"]["mean_lead_hours"] is None
    assert report["summary"]["acceptance_events_in_period"] is None
    assert report["summary"]["acceptances_per_calendar_day"] is None


def test_reported_stage_effort_is_explicit_and_not_silently_discarded():
    history, first, second, performer1, performer2 = populated()
    stage = next(item for item in history.plan.items if item.title == "Stage")
    history.change(15, "assign", stage.id, WorkAssignment(performer1))
    history.change(16, "move_item", stage.id, "running")
    history.change(17, "record_progress", stage.id, 0.5, 2)
    report = history.report(18)
    assert report["summary"]["reported_effort_hours"] == 7
    assert report["summary"]["leaf_reported_effort_hours"] == 5
    assert next(row for row in report["workload"] if row["performer_id"] == performer1)["reported_effort_hours"] == 7
    assert report["items"][0]["changes"][-1]["actor_name"] == "Engineer"
    assert report["forecast"]["scope"] == "filtered_leaf_execution"


def test_compute_reservations_use_real_pipeline_contract_and_remain_distinct_from_utilization():
    from kraken_manager.domain.pipeline import (
        LayerPipelineDefinition,
        PipelineOperation,
        PipelinePolicy,
        PipelineResources,
    )
    from dataclasses import replace

    history = History()
    layer_id = uid()
    operation = PipelineOperation(
        uid(),
        "Vectorize",
        "frames.vectorize.v1",
        "1.0",
        uid(),
        uid(),
        "image",
        "vector",
        ((1, 1),),
        resources=PipelineResources(2, 1, 512),
    )
    definition = LayerPipelineDefinition(
        history.plan.project_id, layer_id, (operation,), PipelinePolicy(2, PipelineResources(4, 2, 1024))
    )
    state = {
        "project_id": history.plan.project_id,
        "layer_id": layer_id,
        "definition": definition.to_dict(),
        "status": "running",
        "nodes": {operation.id: {"status": "running"}},
    }
    running = EventEnvelope.create(
        stream_id=f"layer-pipeline:{layer_id}",
        project_id=history.plan.project_id,
        revision=1,
        event_type="LayerPipelineWorkflowV1",
        payload={"workflow": state},
        actor=ActorSnapshot.from_principal(history.actor),
        recorded_at=BASE + timedelta(hours=1),
    )
    state = state | {"status": "completed", "nodes": {operation.id: {"status": "succeeded"}}}
    completed = replace(
        running, event_id=uid(), revision=2, payload={"workflow": state}, recorded_at=BASE + timedelta(hours=2)
    )
    source = [*history.events, running, completed]
    report = WorkPlanAnalytics.report(
        source, project_id=history.plan.project_id, as_of=BASE + timedelta(hours=1, minutes=30)
    )
    reservation = report["compute_workload"][0]
    assert reservation["reserved"] == {"cpu": 2, "gpu": 1, "memory_mb": 512}
    assert reservation["reservation_fraction"] == {"cpu": 0.5, "gpu": 0.5, "memory_mb": 0.5}
    assert reservation["measured_utilization"] is None
    assert running.event_id in report["source_event_ids"] and completed.event_id not in report["source_event_ids"]
    later = WorkPlanAnalytics.report(source, project_id=history.plan.project_id, as_of=BASE + timedelta(hours=3))
    assert later["compute_workload"][0]["reserved"] == {"cpu": 0, "gpu": 0, "memory_mb": 0}


def test_streaming_and_unordered_analytics_have_equivalent_facts():
    history, first, second, performer1, performer2 = populated()
    arguments = dict(project_id=history.plan.project_id, as_of=BASE + timedelta(hours=15), start=BASE)
    streamed = WorkPlanAnalytics.report(iter(history.events), assume_sorted=True, **arguments)
    unordered = WorkPlanAnalytics.report(reversed(history.events), **arguments)
    assert streamed == unordered
