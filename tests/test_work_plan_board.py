from threading import Event

from PyQt6.QtCore import QSettings, Qt
from PyQt6.QtWidgets import QFileDialog

from kraken_hub.presentation.qt.work_plan_board import WorkPlanBoardDialog, WorkPlanBoardWidget
from kraken_hub.presentation.qt.work_planning_dialog import WorkPlanningDialog
from tests.test_work_plan_analytics_integration import seed_plan
from tests.test_work_planning_integration import setup_project


def wait_board(qtbot, board):
    qtbot.waitUntil(lambda: board._worker is None, timeout=15000)
    assert board.report is not None, board.status.text()


def setup_board(qtbot, tmp_path, language="English"):
    local, principal, project, layer, performer = setup_project(tmp_path)
    plan, first, second, before, revision = seed_plan(local, principal, project, layer, performer)
    settings = QSettings(str(tmp_path / "board.ini"), QSettings.Format.IniFormat)
    widget = WorkPlanBoardWidget(local, principal, project.id, language=language, settings=settings)
    qtbot.addWidget(widget)
    widget.resize(1250, 900)
    widget.show()
    wait_board(qtbot, widget)
    return widget, local, principal, project, layer, performer, plan, first, second, settings


def test_shown_board_has_named_filters_hierarchy_events_metrics_and_task_drilldown(qtbot, tmp_path):
    widget, local, principal, project, layer, performer, plan, first, second, settings = setup_board(qtbot, tmp_path)
    assert widget.isVisible()
    assert widget.columns["accepted"].count() == widget.columns["blocked"].count() == 1
    assert widget.layers.findText(layer.name) >= 0
    assert widget.performers.findText(performer.name) >= 0
    assert widget.hierarchy.topLevelItemCount() == 2
    column = widget.columns["blocked"]
    qtbot.mouseClick(column.viewport(), Qt.MouseButton.LeftButton, pos=column.visualItemRect(column.item(0)).center())
    assert "Awaiting microscope" in widget.details.toPlainText()
    assert performer.name in widget.details.toPlainText()
    assert widget.changes.rowCount() >= 2
    with qtbot.waitSignal(widget.workItemActivated, timeout=2000) as signal:
        qtbot.mouseClick(widget.open_button, Qt.MouseButton.LeftButton)
    assert signal.args == [plan.id, second]
    metric_row = next(
        r
        for r in range(widget.metrics.rowCount())
        if widget.metrics.item(r, 0).data(Qt.ItemDataRole.UserRole) == "readiness_fraction"
    )
    widget.tabs.setCurrentWidget(widget.metrics)
    rect = widget.metrics.visualItemRect(widget.metrics.item(metric_row, 0))
    qtbot.mouseClick(widget.metrics.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
    qtbot.mouseClick(widget.metric_drilldown_button, Qt.MouseButton.LeftButton)
    assert widget.tabs.currentIndex() == 0
    assert widget.columns["accepted"].count() == 1 and widget.columns["blocked"].count() == 0
    qtbot.mouseClick(widget.clear_drilldown, Qt.MouseButton.LeftButton)
    assert widget.columns["blocked"].count() == 1
    widget.states.setCurrentIndex(widget.states.findData("blocked"))
    wait_board(qtbot, widget)
    assert widget.report["state_counts"] == {"blocked": 1}
    assert widget.columns["accepted"].count() == 0
    widget.search.setText("microscope")
    assert widget.columns["blocked"].count() == 1
    widget.close()
    reopened = WorkPlanBoardWidget(local, principal, project.id, settings=settings, language="English")
    qtbot.addWidget(reopened)
    reopened.show()
    wait_board(qtbot, reopened)
    assert reopened.states.currentData() == "blocked"
    assert reopened.search.text() == "microscope"
    assert reopened.report["state_counts"] == {"blocked": 1}
    reopened.close()


def test_board_capacity_and_xlsx_export_use_real_services_in_background(qtbot, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    widget, local, principal, project, layer, performer, plan, first, second, settings = setup_board(qtbot, tmp_path)
    widget.capacity_performer.setCurrentIndex(widget.capacity_performer.findData(performer.id))
    widget.capacity.setValue(6)
    qtbot.mouseClick(widget.set_capacity, Qt.MouseButton.LeftButton)
    wait_board(qtbot, widget)
    assert widget.report["workload"][0]["capacity_hours_per_calendar_day"] == 6
    assert "Unknown" in widget.forecast.text()
    output = tmp_path / "board.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(output), "XLSX"))
    qtbot.mouseClick(widget.xlsx_button, Qt.MouseButton.LeftButton)
    wait_board(qtbot, widget)
    assert output.is_file()
    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert "Baselines" in workbook.sheetnames and "History" in workbook.sheetnames
        assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    finally:
        workbook.close()
    widget.close()


def test_planning_dialog_board_card_opens_exact_existing_task(qtbot, tmp_path):
    local, principal, project, layer, performer = setup_project(tmp_path)
    plan, first, second, before, revision = seed_plan(local, principal, project, layer, performer)
    dialog = WorkPlanningDialog(local, principal, project.id, language="English")
    qtbot.addWidget(dialog)
    dialog.show()
    qtbot.waitUntil(lambda: dialog._worker is None, timeout=10000)
    qtbot.mouseClick(dialog.board_button, Qt.MouseButton.LeftButton)
    board_dialog = dialog._board_dialog
    qtbot.addWidget(board_dialog)
    wait_board(qtbot, board_dialog.board)
    column = board_dialog.board.columns["blocked"]
    qtbot.mouseClick(column.viewport(), Qt.MouseButton.LeftButton, pos=column.visualItemRect(column.item(0)).center())
    qtbot.mouseClick(board_dialog.board.open_button, Qt.MouseButton.LeftButton)
    qtbot.waitUntil(
        lambda: dialog._worker is None and dialog.selected_item() is not None and dialog.selected_item().id == second,
        timeout=10000,
    )
    assert dialog.plan.id == plan.id and dialog.selected_item().state == "blocked"
    board_dialog.close()
    dialog.close()


def test_board_close_waits_for_running_read_and_keeps_worker_registered(qtbot, tmp_path):
    local, principal, project, layer, performer = setup_project(tmp_path)
    entered, release = Event(), Event()

    class SlowService:
        def __getattr__(self, name):
            return getattr(local, name)

        def work_plan_analytics(self, *args, **kwargs):
            entered.set()
            if not release.wait(10):
                raise TimeoutError("Test release was not signalled")
            return local.work_plan_analytics(*args, **kwargs)

    workers = set()
    dialog = WorkPlanBoardDialog(SlowService(), principal, project.id, language="Русский", background_workers=workers)
    qtbot.addWidget(dialog)
    dialog.show()
    qtbot.waitUntil(entered.is_set, timeout=10000)
    assert len(workers) == 1
    dialog.close()
    assert dialog.isVisible()
    release.set()
    wait_board(qtbot, dialog.board)
    assert not workers
    dialog.close()
    assert not dialog.isVisible()


def test_statistics_page_keeps_qt_responsive_during_real_background_report_read(qtbot, tmp_path):
    from PyQt6.QtCore import QTimer
    from PyQt6.QtWidgets import QTabWidget
    from kraken_hub.presentation.qt.parts.statistics_panel import _statistics_panel

    local, principal, project, layer, performer = setup_project(tmp_path)
    entered, release = Event(), Event()

    class SlowService:
        def __getattr__(self, name):
            return getattr(local, name)

        def activity_records(self):
            entered.set()
            if not release.wait(10):
                raise TimeoutError("Test release was not signalled")
            return local.activity_records()

    workers = set()
    panel = _statistics_panel(SlowService(), principal, background_workers=workers, async_reads=True)
    qtbot.addWidget(panel)
    panel.resize(1250, 900)
    panel.show()
    qtbot.waitUntil(entered.is_set, timeout=10000)
    ticks = []
    QTimer.singleShot(0, lambda: ticks.append(True))
    qtbot.waitUntil(lambda: bool(ticks), timeout=1000)
    assert panel._report_worker is not None and workers
    release.set()
    qtbot.waitUntil(lambda: panel._report_worker is None and not workers, timeout=15000)
    tabs = panel.findChild(QTabWidget, "statisticsTabs")
    assert tabs.count() == 6
    assert panel.findChild(WorkPlanBoardWidget) is not None
    panel.close()
